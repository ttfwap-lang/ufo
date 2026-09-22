"""Live showcase runner: task loading, verifiers on generated files, and cleanup safety."""
import zipfile
from unittest import mock

import pytest

from ufo.tests.eval_suite.live import run_live


def test_load_tasks_substitutes_sandbox_and_checks_dependencies(tmp_path):
    tasks = run_live.load_tasks(tmp_path)
    ids = [t.id for t in tasks]
    assert ids[:2] == ["N1", "W1"] and "G1" in ids
    assert all("{sandbox}" not in t.request for t in tasks)
    assert str(tmp_path) in tasks[0].request
    assert {t.verify["type"] for t in tasks} <= set(run_live.VERIFIERS)

    bad = tmp_path / "bad.yaml"
    bad.write_text("tasks:\n  - {id: A, title: a, request: r, verify: {type: fs}, depends_on: [Z]}\n", encoding="utf-8")
    with pytest.raises(ValueError):
        run_live.load_tasks(tmp_path, bad)


def test_file_text_and_regex(tmp_path):
    (tmp_path / "calc.txt").write_text("Result: 69,104\n", encoding="utf-8")
    ok, _ = run_live.v_file_text(tmp_path, {"path": "calc.txt", "contains": ["69104"], "normalize_digits": True})
    assert ok
    ok, detail = run_live.v_file_text(tmp_path, {"path": "calc.txt", "contains": ["70000"]})
    assert not ok and "70000" in detail
    (tmp_path / "python.txt").write_text("3.14.2", encoding="utf-8")
    assert run_live.v_file_regex(tmp_path, {"path": "python.txt", "pattern": r"\b3\.\d+(\.\d+)?\b"})[0]
    assert run_live.v_file_text(tmp_path, {"path": "missing.txt", "contains": []}) == (False, "missing.txt not found")


def test_docx(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("UFO Showcase", level=1)
    p = d.add_paragraph("This report was written by ")
    p.add_run("UFO").bold = True
    p.add_run(".")
    d.save(str(tmp_path / "report.docx"))
    spec = {"path": "report.docx", "heading": "UFO Showcase", "heading_style": "Heading 1",
            "paragraph": "This report was written by", "bold": "UFO"}
    assert run_live.v_docx(tmp_path, spec)[0]
    assert not run_live.v_docx(tmp_path, dict(spec, bold="report"))[0]


def test_xlsx_with_formula_and_chart(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    rows = [("Region", "Sales"), ("North", 120), ("South", 95), ("East", 143), ("West", 88), ("Central", 104)]
    for r in rows:
        ws.append(r)
    ws["B7"] = "=SUM(B2:B6)"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    ws.add_chart(chart, "D2")
    wb.save(str(tmp_path / "sales.xlsx"))
    spec = {"path": "sales.xlsx", "cells": {"A1": "Region", "B2": 120, "B6": 104},
            "formula_cell": "B7", "formula_prefix": "=SUM(B2:B6", "min_charts": 1}
    assert run_live.v_xlsx(tmp_path, spec) == (True, "cells, formula and chart present")
    assert not run_live.v_xlsx(tmp_path, dict(spec, min_charts=2))[0]


def test_pptx(tmp_path):
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    for i, title in enumerate(["UFO Showcase", "What UFO Did", "Next Steps"]):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        if i == 1:
            slide.placeholders[1].text = "Opened apps\nEdited documents\nVerified results"
    prs.save(str(tmp_path / "deck.pptx"))
    spec = {"path": "deck.pptx", "titles": ["UFO Showcase", "What UFO Did", "Next Steps"],
            "text": ["Opened apps", "Verified results"]}
    assert run_live.v_pptx(tmp_path, spec)[0]
    assert not run_live.v_pptx(tmp_path, dict(spec, titles=["Missing"]))[0]


def test_png_not_blank(tmp_path):
    from PIL import Image, ImageDraw

    blank = Image.new("RGB", (200, 100), "white")
    blank.save(tmp_path / "blank.png")
    drawn = blank.copy()
    ImageDraw.Draw(drawn).rectangle([50, 20, 150, 80], fill="blue")
    drawn.save(tmp_path / "drawing.png")
    assert not run_live.v_png_not_blank(tmp_path, {"path": "blank.png"})[0]
    assert run_live.v_png_not_blank(tmp_path, {"path": "drawing.png"})[0]


def test_fs_and_winver(tmp_path):
    (tmp_path / "archive").mkdir()
    (tmp_path / "archive" / "notes_old.txt").write_text("x", encoding="utf-8")
    assert run_live.v_fs(tmp_path, {"exists": ["archive/notes_old.txt"], "missing": ["notes.txt"]})[0]
    version = run_live.windows_display_version()
    (tmp_path / "winver.txt").write_text(f"Windows 11 version {version}", encoding="utf-8")
    assert run_live.v_winver(tmp_path, {"path": "winver.txt"})[0]


def test_dgx_report(tmp_path):
    (tmp_path / "dgx_report.txt").write_text("Free space on /: 2.7T\nRunning containers: 3\n", encoding="utf-8")
    with mock.patch.object(run_live, "dgx_readings", return_value=("2.7T", 3)):
        assert run_live.v_dgx_report(tmp_path, {"path": "dgx_report.txt"})[0]
    with mock.patch.object(run_live, "dgx_readings", return_value=("2.7T", 5)):
        assert not run_live.v_dgx_report(tmp_path, {"path": "dgx_report.txt"})[0]


def test_crashing_verifier_is_inconclusive(tmp_path):
    task = run_live.LiveTask(id="X", title="x", request="r", verify={"type": "file_text", "path": "a.txt"})
    with mock.patch.dict(run_live.VERIFIERS, {"file_text": mock.Mock(side_effect=RuntimeError("boom"))}):
        assert run_live.verify(task, tmp_path) == (None, "verifier error: boom")


def test_cleanup_only_touches_processes_started_by_the_task(tmp_path):
    task = run_live.LiveTask(id="N1", title="n", request="r", verify={}, apps=["notepad.exe", "explorer.exe"])
    before = {100: "notepad.exe", 1: "explorer.exe"}  # the user's own Notepad and the shell
    after = {100: "notepad.exe", 1: "explorer.exe", 200: "notepad.exe"}
    terminated = []

    class FakeProc:
        def __init__(self, pid):
            self.pid = pid

        def terminate(self):
            terminated.append(self.pid)

    with mock.patch.object(run_live, "process_snapshot", return_value=after), \
         mock.patch.object(run_live, "_close_explorer_windows") as close_explorer, \
         mock.patch("win32gui.EnumWindows"), mock.patch("time.sleep"), \
         mock.patch("psutil.Process", side_effect=FakeProc):
        closed = run_live.cleanup(task, before, tmp_path)
    assert terminated == [200]
    assert closed == ["notepad.exe:200"]
    close_explorer.assert_called_once()


def test_report(tmp_path):
    results = [run_live.TaskResult("N1", "Notepad", "pass", "ok", 12.3, [{"task": "t", "success": True}]),
               run_live.TaskResult("F1", "Explorer", "skipped", "depends on ['N1']")]
    path = run_live.write_report(results, tmp_path)
    text = path.read_text(encoding="utf-8")
    assert "1/2 tasks verified" in text and "| N1 Notepad | PASS | 1 |" in text
    assert (tmp_path / "report.json").is_file()
